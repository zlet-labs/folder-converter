import os
import sys
from pathlib import Path
import random
import datetime
from PIL import Image, ImageDraw, ImageFont, ImageFilter
import numpy as np
import fpdf
from docx import Document as DocxDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import docx.opc.constants
import pptx
from pptx.util import Inches, Pt
import openpyxl

FIXTURES_DIR = Path("evaluation/fixtures")
FIXTURES_DIR.mkdir(parents=True, exist_ok=True)
FONTS_DIR = Path(os.environ.get("WINDIR", "C:\\Windows")) / "Fonts"

def get_font_path(font_name):
    p = FONTS_DIR / font_name
    if p.exists():
        return str(p)
    return None

def create_f01_simple_text():
    pdf = fpdf.FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 10, "Introduction to Document Analysis", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(5)
    pdf.set_font("Helvetica", "", 12)
    pdf.multi_cell(0, 7, "Document analysis is a foundational discipline within modern information retrieval and knowledge engineering. It encompasses the systematic extraction of structured and unstructured information from diverse digital document formats. Modern document pipelines require resilient parsers capable of interpreting typographical conventions and semantic hierarchies.")
    pdf.ln(5)
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, "Core Concepts and Principles", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    pdf.set_font("Helvetica", "", 12)
    pdf.multi_cell(0, 7, "The primary challenge in document conversion lies in bridging visual presentation with machine-actionable structure. Visual documents such as PDFs encode absolute glyph placements rather than semantic document objects such as headings, paragraphs, and list items.")
    pdf.ln(5)
    pdf.multi_cell(0, 7, "A robust converter must accurately infer the underlying document graph without hallucinating non-existent formatting or omitting critical textual passages.")
    out_path = FIXTURES_DIR / "F01_simple_text.pdf"
    pdf.output(str(out_path))
    print(f"Created {out_path}")

def create_f02_multicolumn():
    pdf = fpdf.FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "Comparative Multi-Column Layout Analysis", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(5)
    
    col_w = 90
    left_x = 10
    right_x = 110
    top_y = pdf.get_y()
    
    # Challenge stream order: physically interleave drawing operations between columns
    # 1. Left Heading
    pdf.set_xy(left_x, top_y)
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(col_w, 7, "Section A: Primary Findings", new_x="RIGHT", new_y="TOP")
    
    # 2. Right Heading (drawn immediately after left heading in stream)
    pdf.set_xy(right_x, top_y)
    pdf.cell(col_w, 7, "Section B: Secondary Observations", new_x="RIGHT", new_y="TOP")
    
    # 3. Left Column Paragraph 1
    pdf.set_xy(left_x, top_y + 10)
    pdf.set_font("Helvetica", "", 10)
    pdf.multi_cell(col_w, 6, "Left Column Paragraph 1: Academic and periodical publishing frequently employs multi-column grids to optimize line length for human readability. A naive sequential reading order parser will traverse across columns, interweaving disparate thoughts.")
    left_p1_end_y = pdf.get_y()
    
    # 4. Right Column Paragraph 1 (drawn immediately after left paragraph 1 in stream)
    pdf.set_xy(right_x, top_y + 10)
    pdf.multi_cell(col_w, 6, "Right Column Paragraph 1: Reading order verification is essential when validating scientific papers. Textual continuity must be strictly preserved across column breaks.")
    right_p1_end_y = pdf.get_y()
    
    # 5. Left Column Paragraph 2
    pdf.set_xy(left_x, left_p1_end_y + 3)
    pdf.multi_cell(col_w, 6, "Left Column Paragraph 2: High-fidelity layout models identify physical column boundaries and order document elements down each column in sequence before advancing to the subsequent column.")
    
    # 6. Right Column Paragraph 2
    pdf.set_xy(right_x, right_p1_end_y + 3)
    pdf.multi_cell(col_w, 6, "Right Column Paragraph 2: If this paragraph appears in markdown immediately following Right Column Paragraph 1, the column layout analyzer is performing correctly.")
    
    out_path = FIXTURES_DIR / "F02_multicolumn.pdf"
    pdf.output(str(out_path))
    print(f"Created {out_path} (with interleaved stream order to challenge 2D layout reading order)")

def create_f03_table():
    pdf = fpdf.FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "Quarterly Performance Table", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(5)
    
    headers = ["Item ID", "Product Category", "Units Sold", "Total Revenue"]
    col_widths = [30, 60, 40, 50]
    data = [
        ["ITM-001", "Desktop Application", "1250", "$62,500.00"],
        ["ITM-002", "Cloud Worker Engine", "840", "$105,000.00"],
        ["ITM-003", "Enterprise Support", "45", "$33,750.00"],
        ["ITM-004", "Developer Tooling", "310", "$15,500.00"],
    ]
    
    pdf.set_font("Helvetica", "B", 11)
    for w, h in zip(col_widths, headers):
        pdf.cell(w, 8, h, border=1)
    pdf.ln()
    
    pdf.set_font("Helvetica", "", 10)
    for row in data:
        for w, val in zip(col_widths, row):
            pdf.cell(w, 8, val, border=1)
        pdf.ln()
        
    out_path = FIXTURES_DIR / "F03_table.pdf"
    pdf.output(str(out_path))
    print(f"Created {out_path}")

def create_f04_header_footer():
    class HeaderFooterPDF(fpdf.FPDF):
        def header(self):
            self.set_font("Helvetica", "I", 9)
            self.cell(0, 8, "Company Confidential Report - Zlet Evaluation Suite", border="B", align="L")
            self.ln(10)
        def footer(self):
            self.set_y(-15)
            self.set_font("Helvetica", "I", 9)
            self.cell(0, 10, f"Page {self.page_no()} of {{nb}}", border="T", align="C")

    pdf = HeaderFooterPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, "Executive Overview: Section One", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 11)
    pdf.multi_cell(0, 7, "This is the primary body of page one. Repeated headers and footers contain metadata that can pollute extracted text if not correctly segmented as non-body artifacts.")
    
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, "Executive Overview: Section Two", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 11)
    pdf.multi_cell(0, 7, "This is the body of page two. Evaluators must verify whether 'Company Confidential Report' or 'Page 2 of 2' are improperly treated as headings or body text.")
    
    out_path = FIXTURES_DIR / "F04_header_footer.pdf"
    pdf.output(str(out_path))
    print(f"Created {out_path}")

def create_f05_broken_wrap():
    pdf = fpdf.FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, "Hyphenation and Line Wrapping Test", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(5)
    pdf.set_font("Helvetica", "", 11)
    
    lines = [
        "Software engineering methodologies frequently encoun-",
        "ter challenges when reconciling complex architectural con-",
        "straints with aggressive delivery schedules. In docu-",
        "ment parsing, dehyphenation algorithms must reassemble bro-",
        "ken fragments into coherent tokens without corrupting genu-",
        "ine compound nouns such as state-of-the-art implementations."
    ]
    for line in lines:
        pdf.cell(0, 6, line, new_x="LMARGIN", new_y="NEXT")
        
    out_path = FIXTURES_DIR / "F05_broken_wrap.pdf"
    pdf.output(str(out_path))
    print(f"Created {out_path}")

# Ground truth text for F06 scanned OCR reference validation
F06_TITLE = "OFFICIAL NOTICE: OCR CAPABILITY TEST"
F06_BODY = (
    "Document Identifier: SC-98421\n"
    "Date of Certification: September 2026\n"
    "Issuer: Zlet Systems Quality Assessment Group\n\n"
    "This scanned image tests whether the OCR subsystem activates\n"
    "accurately on non-searchable rasterized PDF pages with synthetic scan degradation.\n"
    "All characters in this paragraph must be recognized without errors.\n"
    "The quick brown fox jumps over the lazy dog.\n"
    "Expected result: clean extracted markdown text with intact wording."
)

def create_f06_scanned():
    # Representative scanned-document fixture with realistic degradation
    w, h = 1600, 1200
    # 1. Warm off-white background simulating physical scanner paper
    img = Image.new("RGB", (w, h), color=(247, 245, 240))
    draw = ImageDraw.Draw(img)
    
    font_path = get_font_path("arial.ttf")
    font_title = ImageFont.truetype(font_path, 44) if font_path else ImageFont.load_default()
    font_body = ImageFont.truetype(font_path, 30) if font_path else ImageFont.load_default()
    
    # 2. Dark gray scanner ink
    ink_color = (35, 35, 35)
    draw.text((120, 110), F06_TITLE, fill=ink_color, font=font_title)
    draw.line([(120, 170), (1480, 170)], fill=ink_color, width=3)
    
    draw.text((120, 210), F06_BODY, fill=ink_color, font=font_body, spacing=14)
    
    # 3. Add scanner noise using deterministic seeded generator
    # Fixed seed 42 guarantees byte-reproducible synthetic scanner noise
    rng = np.random.default_rng(42)
    img_arr = np.array(img, dtype=np.int16)
    noise = rng.normal(0, 3.5, img_arr.shape).astype(np.int16)
    noisy_arr = np.clip(img_arr + noise, 0, 255).astype(np.uint8)
    img = Image.fromarray(noisy_arr)
    
    # 4. Add subtle optical blur
    img = img.filter(ImageFilter.GaussianBlur(radius=0.5))
    
    # 5. Add rotational scanner skew (-0.75 degrees)
    img = img.rotate(-0.75, resample=Image.Resampling.BICUBIC, expand=False, fillcolor=(247, 245, 240))
    
    img_tmp_path = FIXTURES_DIR / "f06_temp.png"
    img.save(str(img_tmp_path), "PNG")
    
    # Embed into image-only PDF
    pdf = fpdf.FPDF(orientation="P", unit="mm", format="A4")
    # Fixed creation date ensures byte-reproducible PDF metadata across regeneration passes
    pdf.set_creation_date(datetime.datetime(2026, 9, 10, 0, 0, 0, tzinfo=datetime.timezone.utc))
    pdf.add_page()
    pdf.image(str(img_tmp_path), x=10, y=10, w=190)
    out_path = FIXTURES_DIR / "F06_scanned.pdf"
    pdf.output(str(out_path))
    if img_tmp_path.exists():
        img_tmp_path.unlink()
    print(f"Created {out_path} (with synthetic scan degradation: noise, blur, and skew; RNG seed=42)")

def create_f07_cyrillic():
    pdf = fpdf.FPDF()
    pdf.add_page()
    arial_path = get_font_path("arial.ttf")
    if not arial_path:
        raise RuntimeError("arial.ttf required for Cyrillic fixture")
    pdf.add_font("ArialCyr", "", arial_path)
    pdf.add_font("ArialCyr", "B", get_font_path("arialbd.ttf") or arial_path)
    
    pdf.set_font("ArialCyr", "B", 16)
    pdf.cell(0, 10, "Оценка качества конвертации документов", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(5)
    
    pdf.set_font("ArialCyr", "B", 13)
    pdf.cell(0, 8, "Архитектурные требования и локализация", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    
    pdf.set_font("ArialCyr", "", 11)
    text1 = (
        "Обработка русскоязычного текста требует корректной поддержки кодировки UTF-8 "
        "и правильного сопоставления глифов. Любые ошибки в декодировании шрифтовых таблиц "
        "приводят к искажению текста (кракозябрам)."
    )
    pdf.multi_cell(0, 7, text1)
    pdf.ln(4)
    
    text2 = (
        "В десктопном приложении Zlet Converter конвертация документов на русском языке "
        "является критически важным сценарием использования. Модель разметки Docling "
        "должна корректно сохранять заголовки, списки и табличные структуры с кириллицей."
    )
    pdf.multi_cell(0, 7, text2)
    
    out_path = FIXTURES_DIR / "F07_cyrillic.pdf"
    pdf.output(str(out_path))
    print(f"Created {out_path}")

def add_docx_hyperlink(paragraph, url, text, color="0000FF", underline=True):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    if color:
        c = OxmlElement('w:color')
        c.set(qn('w:val'), color)
        rPr.append(c)
    if underline:
        u = OxmlElement('w:u')
        u.set(qn('w:val'), 'single')
        rPr.append(u)
    new_run.append(rPr)
    text_elem = OxmlElement('w:t')
    text_elem.text = text
    new_run.append(text_elem)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

def create_f08_structured_docx():
    doc = DocxDocument()
    doc.add_heading("Zlet Converter Architecture Review", level=1)
    doc.add_paragraph("This document provides a comprehensive structural specification of the conversion pipeline.")
    
    doc.add_heading("Core Functional Capabilities", level=2)
    doc.add_paragraph("Key requirements are tracked in the following list:")
    
    bullet_items = [
        "High-fidelity document layout reconstruction",
        "Deterministic local markdown export",
        "Graceful offline error isolation"
    ]
    for b in bullet_items:
        doc.add_paragraph(b, style="List Bullet")
        
    doc.add_heading("Execution Stages", level=2)
    numbered_items = [
        "Ingest and validate source payload",
        "Parse structural elements into normalized tree",
        "Serialize tree to canonical GitHub-flavored Markdown"
    ]
    for n in numbered_items:
        doc.add_paragraph(n, style="List Number")
        
    doc.add_heading("Capability Matrix", level=2)
    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Format"
    hdr_cells[1].text = "Support Tier"
    hdr_cells[2].text = "Target Status"
    
    r1 = table.rows[1].cells
    r1[0].text = "DOCX"
    r1[1].text = "Primary"
    r1[2].text = "Full Markdown"
    
    r2 = table.rows[2].cells
    r2[0].text = "PDF"
    r2[1].text = "Evaluation"
    r2[2].text = "Under Spike"
    
    # Add real w:hyperlink element with distinct anchor text
    doc.add_paragraph()
    p_link = doc.add_paragraph("For complete source code and updates, visit the ")
    add_docx_hyperlink(p_link, "https://github.com/zlet-labs/zlet-converter", "Zlet Converter GitHub Repository")
    p_link.add_run(".")
    
    out_path = FIXTURES_DIR / "F08_structured.docx"
    doc.save(str(out_path))
    print(f"Created {out_path} (with native w:hyperlink XML element and distinct anchor text)")

def create_f09_slides():
    prs = pptx.Presentation()
    
    slide_layout = prs.slide_layouts[0]
    slide1 = prs.slides.add_slide(slide_layout)
    slide1.shapes.title.text = "Zlet Converter Strategy"
    slide1.placeholders[1].text = "Docling Markdown Conversion Spike (ZC-042)"
    
    slide_layout = prs.slide_layouts[1]
    slide2 = prs.slides.add_slide(slide_layout)
    slide2.shapes.title.text = "Key Evaluation Criteria"
    tf = slide2.placeholders[1].text_frame
    tf.text = "Local execution without cloud dependency"
    p = tf.add_paragraph()
    p.text = "Reading order integrity across complex columns"
    p = tf.add_paragraph()
    p.text = "Precise table and list syntax preservation"
    
    slide_layout = prs.slide_layouts[5]
    slide3 = prs.slides.add_slide(slide_layout)
    slide3.shapes.title.text = "Architecture Milestone Summary"
    
    rows, cols = 3, 3
    table_shape = slide3.shapes.add_table(rows, cols, Inches(1), Inches(2), Inches(8), Inches(2.5))
    table = table_shape.table
    table.cell(0, 0).text = "Milestone"
    table.cell(0, 1).text = "Target Engine"
    table.cell(0, 2).text = "Completion"
    
    table.cell(1, 0).text = "ZC-041"
    table.cell(1, 1).text = "Repo Docs Alignment"
    table.cell(1, 2).text = "Complete"
    
    table.cell(2, 0).text = "ZC-042"
    table.cell(2, 1).text = "Docling Evaluation"
    table.cell(2, 2).text = "In Progress"
    
    out_path = FIXTURES_DIR / "F09_slides.pptx"
    prs.save(str(out_path))
    print(f"Created {out_path}")

def create_f10_sheets():
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Transactions"
    ws1.append(["Transaction ID", "Merchant", "Amount", "Currency", "Status"])
    ws1.append(["TX-1001", "Acme Cloud Services", 450.00, "USD", "Settled"])
    ws1.append(["TX-1002", "Developer Hardware Inc", 1299.99, "USD", "Settled"])
    ws1.append(["TX-1003", "Software Licensing Ltd", 89.00, "EUR", "Pending"])
    ws1.append(["TX-1004", "Network Connectivity Corp", 215.50, "USD", "Settled"])
    
    ws2 = wb.create_sheet(title="SummaryMetrics")
    ws2.append(["Metric Name", "Metric Value"])
    ws2.append(["Total Transactions", 4])
    ws2.append(["Settled Count", 3])
    ws2.append(["Primary Currency", "USD"])
    
    out_path = FIXTURES_DIR / "F10_sheets.xlsx"
    wb.save(str(out_path))
    print(f"Created {out_path}")

def create_f11_html():
    html_content = """<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <title>Structural Document Representation</title>
</head>
<body>
    <h1>Document Conversion Architecture</h1>
    <p>This web document serves as a benchmark for HTML structural interpretation into Markdown format.</p>
    
    <h2>Engine Requirements</h2>
    <ul>
        <li>Preserve heading hierarchies from h1 to h6</li>
        <li>Format ordered and unordered lists accurately</li>
        <li>Convert HTML tables into standard GFM tables</li>
    </ul>
    
    <h2>Performance Overview</h2>
    <table border="1">
        <thead>
            <tr><th>Component</th><th>Language</th><th>Status</th></tr>
        </thead>
        <tbody>
            <tr><td>Desktop UI</td><td>C# / WPF</td><td>Production</td></tr>
            <tr><td>Worker Host</td><td>Python 3.11</td><td>Evaluation</td></tr>
        </tbody>
    </table>
    
    <p>Project documentation is hosted at <a href="https://github.com/zlet-labs/zlet-converter">Zlet Converter Repository</a>.</p>
</body>
</html>"""
    out_path = FIXTURES_DIR / "F11_structural.html"
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(html_content)
    print(f"Created {out_path}")

def create_f12_nested_lists():
    pdf = fpdf.FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "Hierarchical Nested List Document", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(5)
    
    pdf.set_font("Helvetica", "", 11)
    items = [
        "1. Top Level Requirement: Robust Pipeline",
        "    a. Sub-requirement: Parse PDF structures",
        "        i. Validate heading levels",
        "        ii. Validate table cells",
        "    b. Sub-requirement: Parse DOCX structures",
        "2. Top Level Requirement: Isolated Runtime",
        "    a. Sub-requirement: Stdin/stdout process communication",
        "    b. Sub-requirement: Process cancellation on user request",
        "3. Top Level Requirement: Zero Data Leakage",
        "    * Fully offline model execution",
        "    * No external telemetry reporting"
    ]
    for item in items:
        pdf.cell(0, 7, item, new_x="LMARGIN", new_y="NEXT")
        
    out_path = FIXTURES_DIR / "F12_nested_lists.pdf"
    pdf.output(str(out_path))
    print(f"Created {out_path}")

def create_f13_empty():
    pdf = fpdf.FPDF()
    pdf.add_page()
    out_path = FIXTURES_DIR / "F13_empty.pdf"
    pdf.output(str(out_path))
    print(f"Created {out_path}")

def create_f14_corrupted():
    out_path = FIXTURES_DIR / "F14_corrupted.pdf"
    with open(out_path, "wb") as f:
        f.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
        f.write(b"CORRUPTED STREAM CONTENT NON RECOVERABLE BYTES 0xDEADBEEF\n")
        f.write(b"1 0 obj << /Type /Catalog >> endobj\n")
    print(f"Created {out_path}")

def create_f15_mixed_unicode():
    pdf = fpdf.FPDF()
    pdf.add_page()
    arial_path = get_font_path("arial.ttf")
    
    cjk_fonts = ["simsun.ttc", "msyh.ttc", "yugothm.ttc", "meiryo.ttc", "msgothic.ttc"]
    cjk_path = None
    for f in cjk_fonts:
        p = get_font_path(f)
        if p:
            cjk_path = p
            break
            
    if not cjk_path:
        raise RuntimeError("No CJK TrueType font found on system (checked simsun.ttc, msyh.ttc, etc.). CJK font is strictly required for F15.")
        
    pdf.add_font("ArialUni", "", arial_path)
    pdf.add_font("ArialUni", "B", get_font_path("arialbd.ttf") or arial_path)
    
    pdf.set_font("ArialUni", "B", 16)
    pdf.cell(0, 10, "Multilingual Unicode Verification Suite", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(5)
    
    pdf.set_font("ArialUni", "B", 12)
    pdf.cell(0, 8, "1. Latin Script Segment", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("ArialUni", "", 11)
    pdf.multi_cell(0, 6, "Standard Latin English text with diacritics: café, façade, über, résumé.")
    pdf.ln(4)
    
    pdf.set_font("ArialUni", "B", 12)
    pdf.cell(0, 8, "2. Cyrillic Script Segment (Кириллица)", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("ArialUni", "", 11)
    pdf.multi_cell(0, 6, "Тестирование кириллического алфавита: Привет мир, алгоритмы обработки данных, надежность.")
    pdf.ln(4)
    
    # CJK segment is required - fail if not present
    pdf.add_font("CJKFont", "", cjk_path)
    pdf.set_font("CJKFont", "", 12)
    pdf.cell(0, 8, "3. CJK Script Segment (中日韩文字)", new_x="LMARGIN", new_y="NEXT")
    pdf.multi_cell(0, 7, "测试中文文档解析能力: 你好世界。文档转换质量必须保持字形和语义的完整。")
    
    out_path = FIXTURES_DIR / "F15_mixed_unicode.pdf"
    pdf.output(str(out_path))
    print(f"Created {out_path}")

def main():
    print("Generating evaluation fixtures...")
    create_f01_simple_text()
    create_f02_multicolumn()
    create_f03_table()
    create_f04_header_footer()
    create_f05_broken_wrap()
    create_f06_scanned()
    create_f07_cyrillic()
    create_f08_structured_docx()
    create_f09_slides()
    create_f10_sheets()
    create_f11_html()
    create_f12_nested_lists()
    create_f13_empty()
    create_f14_corrupted()
    create_f15_mixed_unicode()
    print("All fixtures generated successfully.")

if __name__ == "__main__":
    main()

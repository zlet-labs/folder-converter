#!/usr/bin/env python3
"""
zlet_docling_worker.py - Out-of-process Document-to-Markdown worker for Zlet Converter.
Communicates via line-delimited JSON over stdin / stdout.
"""

import sys
import os
import json
import tempfile
import datetime

# Ensure strict offline execution and suppress benign symlink warnings
os.environ["HF_HUB_OFFLINE"] = "1"
os.environ["TRANSFORMERS_OFFLINE"] = "1"
os.environ["HF_HUB_DISABLE_SYMLINKS_WARNING"] = "1"
os.environ["PYTHONIOENCODING"] = "utf-8"
os.environ["PYTHONUTF8"] = "1"

_converter = None


def validate_environment():
    """Validates Python version and critical dependencies."""
    if not (sys.version_info >= (3, 10) and sys.version_info < (3, 13)):
        return False, f"Unsupported Python version: {sys.version.split()[0]} (expected 3.10-3.12)"
    try:
        import docling
        import openpyxl
        import bs4
    except ImportError as e:
        return False, f"Missing required dependency: {e.name}"
    return True, ""


def get_converter():
    global _converter
    if _converter is None:
        from docling.document_converter import DocumentConverter, PdfFormatOption
        from docling.datamodel.pipeline_options import PdfPipelineOptions

        # Production baseline: do NOT enable OCR. Scanned/image-only PDFs must fail gracefully.
        pdf_options = PdfPipelineOptions(do_ocr=False)
        _converter = DocumentConverter(
            format_options={"pdf": PdfFormatOption(pipeline_options=pdf_options)}
        )
    return _converter


def convert_pdf(source_path):
    if os.path.getsize(source_path) == 0:
        return ""
    converter = get_converter()
    res = converter.convert(source_path)
    md = res.document.export_to_markdown()
    if not md or not md.strip():
        raise ValueError("scanned_pdf_unsupported: PDF contains no extractable text.")
    return md


def convert_docx(source_path):
    converter = get_converter()
    res = converter.convert(source_path)
    md = res.document.export_to_markdown()
    return md


def convert_pptx(source_path):
    converter = get_converter()
    res = converter.convert(source_path)
    md = res.document.export_to_markdown()
    return md


def format_cell_value(cell):
    val = cell.value
    if val is None:
        return ""

    if isinstance(val, str):
        # Never output raw formula
        if val.startswith("="):
            return ""
        return val.strip()

    # Dates
    if getattr(cell, "is_date", False) or isinstance(val, (datetime.date, datetime.datetime, datetime.time)):
        if isinstance(val, datetime.datetime):
            if val.hour == 0 and val.minute == 0 and val.second == 0 and val.microsecond == 0:
                return val.strftime("%Y-%m-%d")
            return val.strftime("%Y-%m-%d %H:%M:%S")
        elif isinstance(val, datetime.date):
            return val.strftime("%Y-%m-%d")
        elif isinstance(val, datetime.time):
            return val.strftime("%H:%M:%S")

    fmt = getattr(cell, "number_format", "") or ""

    if isinstance(val, (int, float)):
        if "%" in fmt:
            dec_places = 2
            if "0.00%" in fmt:
                dec_places = 2
            elif "0.0%" in fmt:
                dec_places = 1
            elif "0%" in fmt:
                dec_places = 0
            return f"{val * 100:.{dec_places}f}%"

        for sym in ("$", "€", "£", "¥", "₽"):
            if sym in fmt:
                if isinstance(val, float) and val.is_integer():
                    return f"{sym}{int(val)}"
                elif isinstance(val, (int, float)):
                    return f"{sym}{val:,.2f}"

        if isinstance(val, float):
            if val.is_integer():
                return str(int(val))
            return f"{val:g}"
        return str(val)

    return str(val).strip()


def convert_xlsx(source_path):
    import openpyxl

    wb = openpyxl.load_workbook(source_path, data_only=True)
    parts = []
    for name in wb.sheetnames:
        ws = wb[name]
        parts.append(f"## {name}\n")
        rows = list(ws.iter_rows())
        content_rows = []
        for r in rows:
            formatted_row = [format_cell_value(c) for c in r]
            if any(c != "" for c in formatted_row):
                content_rows.append(formatted_row)
        if not content_rows:
            parts.append("*(Пустой лист)*\n")
            continue
        max_cols = max(len(r) for r in content_rows)
        padded = [r + [""] * (max_cols - len(r)) for r in content_rows]
        header = padded[0]
        header_str = "| " + " | ".join(c.replace("|", "\\|").replace("\n", " ") for c in header) + " |"
        sep_str = "| " + " | ".join("---" for _ in header) + " |"
        data_strs = [
            "| " + " | ".join(c.replace("|", "\\|").replace("\n", " ") for c in r) + " |"
            for r in padded[1:]
        ]
        parts.append("\n".join([header_str, sep_str] + data_strs) + "\n")
    return "\n\n".join(parts)


def convert_html(source_path):
    from bs4 import BeautifulSoup

    with open(source_path, "rb") as f:
        raw_bytes = f.read()

    html_text = ""
    if raw_bytes.startswith(b"\xef\xbb\xbf"):
        html_text = raw_bytes[3:].decode("utf-8", errors="replace")
    elif raw_bytes.startswith(b"\xff\xfe"):
        html_text = raw_bytes[2:].decode("utf-16-le", errors="replace")
    elif raw_bytes.startswith(b"\xfe\xff"):
        html_text = raw_bytes[2:].decode("utf-16-be", errors="replace")
    else:
        try:
            html_text = raw_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                html_text = raw_bytes.decode("windows-1251")
            except UnicodeDecodeError:
                html_text = raw_bytes.decode("latin-1")

    soup = BeautifulSoup(html_text, "html.parser")
    for tag_name in ("script", "style", "link", "meta", "head"):
        for tag in soup.find_all(tag_name):
            tag.decompose()

    # Strip remote images to prevent SSRF / tracking / network leaks while preserving hyperlinks
    for img in soup.find_all("img"):
        src = img.get("src", "").strip()
        if src.startswith(("http://", "https://", "//", "ftp://")) or "://" in src:
            img.decompose()

    sanitized_html = str(soup)

    # Convert sanitized HTML via DocumentConverter using a temporary file
    temp_fd, temp_file = tempfile.mkstemp(suffix=".html")
    try:
        with os.fdopen(temp_fd, "w", encoding="utf-8") as tf:
            tf.write(sanitized_html)
        converter = get_converter()
        res = converter.convert(temp_file)
        md = res.document.export_to_markdown()
        return md
    finally:
        if os.path.exists(temp_file):
            try:
                os.remove(temp_file)
            except OSError:
                pass


def convert_txt(source_path):
    with open(source_path, "rb") as f:
        raw_bytes = f.read()

    if not raw_bytes:
        return ""

    # Explicit BOM handling followed by strict UTF-8, Windows-1251, and Latin-1 fallback
    text = None
    if raw_bytes.startswith(b"\xef\xbb\xbf"):
        text = raw_bytes[3:].decode("utf-8", errors="replace")
    elif raw_bytes.startswith(b"\xff\xfe"):
        text = raw_bytes[2:].decode("utf-16-le", errors="replace")
    elif raw_bytes.startswith(b"\xfe\xff"):
        text = raw_bytes[2:].decode("utf-16-be", errors="replace")
    else:
        try:
            text = raw_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = raw_bytes.decode("windows-1251")
            except UnicodeDecodeError:
                text = raw_bytes.decode("latin-1")

    # Normalize newlines while preserving paragraph breaks
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    return text


def sanitize_error_message(msg):
    """Sanitizes error messages to remove user home directory and private path leaks."""
    if not msg:
        return ""
    user_home = os.path.expanduser("~")
    if user_home and user_home in msg:
        msg = msg.replace(user_home, "<user_home>")
    return msg


def process_request(req):
    req_id = req.get("id", "")
    source_path = req.get("sourcePath", "")
    output_path = req.get("outputPath", "")
    source_format = req.get("sourceFormat", "").lower()

    if not source_path or not os.path.exists(source_path):
        return {
            "id": req_id,
            "success": False,
            "errorCode": "source_unreadable",
            "errorMessage": "Source file does not exist.",
        }

    try:
        if source_format == "pdf":
            md = convert_pdf(source_path)
        elif source_format == "docx":
            md = convert_docx(source_path)
        elif source_format == "pptx":
            md = convert_pptx(source_path)
        elif source_format == "xlsx":
            md = convert_xlsx(source_path)
        elif source_format in ("html", "htm"):
            md = convert_html(source_path)
        elif source_format == "txt":
            md = convert_txt(source_path)
        else:
            return {
                "id": req_id,
                "success": False,
                "errorCode": "unsupported_format",
                "errorMessage": f"Unsupported format: {source_format}",
            }

        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        with open(output_path, "w", encoding="utf-8", newline="\n") as f:
            f.write(md)

        return {
            "id": req_id,
            "success": True,
            "errorCode": "",
            "errorMessage": "",
        }
    except ValueError as ve:
        err_str = str(ve)
        if "scanned_pdf_unsupported" in err_str:
            return {
                "id": req_id,
                "success": False,
                "errorCode": "scanned_pdf_unsupported",
                "errorMessage": "PDF contains no extractable text.",
            }
        return {
            "id": req_id,
            "success": False,
            "errorCode": "conversion_error",
            "errorMessage": sanitize_error_message(err_str),
        }
    except Exception as ex:
        err_msg = f"{type(ex).__name__}: {str(ex)}"
        return {
            "id": req_id,
            "success": False,
            "errorCode": "docling_conversion_failed",
            "errorMessage": sanitize_error_message(err_msg),
        }


def main():
    env_ok, env_err = validate_environment()

    docling_version = "2.126.0"
    if env_ok:
        try:
            import importlib.metadata
            docling_version = importlib.metadata.version("docling")
        except Exception:
            pass

    greeting = {
        "ready": env_ok,
        "version": "1.0.0",
        "pythonVersion": sys.version.split()[0],
        "doclingVersion": docling_version,
    }
    if not env_ok:
        greeting["errorCode"] = "docling_version_incompatible"
        greeting["errorMessage"] = env_err

    sys.stdout.write(json.dumps(greeting) + "\n")
    sys.stdout.flush()

    for line in sys.stdin:
        line = line.strip()
        if not line:
            continue
        try:
            req = json.loads(line)
            if req.get("command") == "ping":
                resp = {"id": req.get("id", ""), "success": True}
            elif req.get("command") == "shutdown":
                break
            else:
                resp = process_request(req)
        except Exception as e:
            resp = {
                "id": "",
                "success": False,
                "errorCode": "protocol_error",
                "errorMessage": sanitize_error_message(str(e)),
            }
        sys.stdout.write(json.dumps(resp) + "\n")
        sys.stdout.flush()


if __name__ == "__main__":
    main()
